"""reports/new_walkforward_<date>.xlsx + charts (development + walk-forward only)."""
import os
import pickle
import subprocess
import sys
import xml.etree.ElementTree as ET

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from new_research.run_all import load, CFG, WF0          # noqa: E402

DATE = '20261018'
CH = os.path.join(ROOT, 'reports', 'charts')
os.makedirs(CH, exist_ok=True)
PR, MO, SI, BT, EV, LP = (load(x) for x in ('profiles', 'models', 'similarity', 'backtest_up', 'evaluate', 'legacy_picks'))
LR = pickle.load(open(os.path.join(ROOT, 'common_eval', '.cache', 'long_results.pkl'), 'rb'))
BR = pickle.load(open(os.path.join(ROOT, 'common_eval', '.cache', 'bench_results.pkl'), 'rb'))


def rebased(eq):
    e = eq[eq.index >= WF0]
    return e / e.iloc[0]


# ---------------- charts ----------------
def chart_clusters():
    cz = PR['clusters']['centroids_z']
    fig, ax = plt.subplots(figsize=(12, 5))
    w = 0.8 / len(cz)
    x = np.arange(cz.shape[1])
    for c in range(len(cz)):
        ax.bar(x + c * w, cz.iloc[c].values, w, label=f"cluster {c} ({int(PR['clusters']['centroids']['n_events'].iloc[c])} events)")
    ax.set_xticks(x + w); ax.set_xticklabels(cz.columns, rotation=45, ha='right', fontsize=8)
    ax.axhline(0, color='k', lw=.8); ax.set_ylabel('z-score vs all development coin-days')
    ax.set_title(f"Pre-move profile of UP events (k-means, k={PR['clusters']['k']}, silhouette {PR['clusters']['silhouette'][PR['clusters']['k']]:.2f})")
    ax.legend(); plt.tight_layout(); plt.savefig(os.path.join(CH, f'new_cluster_profiles_{DATE}.png'), dpi=120); plt.close()


def chart_equity(mode):
    fig, ax = plt.subplots(2, 1, figsize=(12, 9), sharex=True, gridspec_kw={'height_ratios': [2, 1]})
    series = [(k, BT[k][mode]['eq']) for k in ('NEW_XGB', 'NEW_XGB_regime', 'NEW_BASELINE', 'NEW_BASELINE_regime', 'NEW_LOGISTIC')]
    series += [('LEGACY_FIXED long', LR['L4_+M4_keep_NaN'][mode]['eq']), ('LEGACY_AS_IS long', LR['L0_AS_IS'][mode]['eq'])]
    for n, e in series:
        e = rebased(e)
        ls = '-' if n.startswith('NEW') else ':'
        ax[0].plot(e.index, e.values, ls, label=n, lw=1.3)
        ax[1].plot(e.index, (e / e.cummax() - 1).values, ls, lw=1)
    for n in ('BTC_HOLD', 'EQUAL_WEIGHT'):
        e = rebased(BR[n]['eq'])
        ax[0].plot(e.index, e.values, '--', label=n, color='k' if n == 'BTC_HOLD' else 'gray', lw=1.4)
        ax[1].plot(e.index, (e / e.cummax() - 1).values, '--', color='k' if n == 'BTC_HOLD' else 'gray', lw=1)
    ax[0].set_yscale('log'); ax[0].legend(fontsize=8); ax[0].grid(alpha=.3)
    ax[0].set_title(f'NEW (UP) vs LEGACY vs benchmarks, walk-forward 2022-01 -> 2025-06, exit = {mode}')
    ax[1].set_title('Drawdown'); ax[1].grid(alpha=.3)
    plt.tight_layout(); plt.savefig(os.path.join(CH, f'new_equity_drawdown_{mode}_{DATE}.png'), dpi=120); plt.close()


def chart_random():
    rnd = dict(load('random_up')); rnd.update(load('random_up_NEW_LOGISTIC'))
    fig, axs = plt.subplots(2, 4, figsize=(18, 8))
    for row, key in enumerate(('NEW_XGB', 'NEW_LOGISTIC')):
        for a, (mode, ex) in zip(axs[row], [('timing', 'native'), ('timing', 'unified'), ('rate', 'native'), ('rate', 'unified')]):
            rs = rnd[key][(mode, ex)]['expectancy']
            a.hist(rs, bins=40, color='lightgray', edgecolor='gray')
            v = BT[key][ex]['wf']['expectancy']
            a.axvline(v, color='red', lw=2, label=f'{key} {v:.3f}')
            a.axvline(np.percentile(rs, 95), color='blue', ls='--', label='P95')
            a.set_title(f'{key} expectancy | RANDOM {mode}-matched | {ex}', fontsize=8); a.legend(fontsize=7)
    plt.tight_layout(); plt.savefig(os.path.join(CH, f'new_random_hist_{DATE}.png'), dpi=110); plt.close()


chart_clusters()
for m in ('native', 'unified'):
    chart_equity(m)
chart_random()

# ---------------- unit tests (junit) ----------------
junit = os.path.join(ROOT, 'new_research', '.cache', 'junit.xml')
subprocess.run([sys.executable, '-m', 'pytest', 'tests', '-q', f'--junitxml={junit}'], cwd=ROOT, capture_output=True)
ut = []
for tc in ET.parse(junit).getroot().iter('testcase'):
    st = 'failed' if tc.find('failure') is not None else ('skipped' if tc.find('skipped') is not None else 'passed')
    ut.append({'suite': 'tests/test_guardrails.py (applied)', 'test': tc.get('name'), 'status': st, 'seconds': float(tc.get('time'))})
patched = r'C:/Users/User/AppData/Local/Temp/claude/c--Project-backtest/c8133345-49e8-4a78-8722-bdcdd17b9945/scratchpad/patched'
j2 = os.path.join(ROOT, 'new_research', '.cache', 'junit_funding.xml')
subprocess.run([sys.executable, '-m', 'pytest', 'tests', '-q', f'--junitxml={j2}'], cwd=patched, capture_output=True)
for tc in ET.parse(j2).getroot().iter('testcase'):
    st = 'failed' if tc.find('failure') is not None else 'passed'
    ut.append({'suite': 'proposals/common_eval_funding.diff (NOT applied, awaiting approval)', 'test': tc.get('name'), 'status': st,
               'seconds': float(tc.get('time'))})
UT = pd.DataFrame(ut)

# ---------------- workbook ----------------
wb = Workbook(); wb.remove(wb.active)


def sheet(name, df, index=False):
    ws = wb.create_sheet(name[:31])
    d = df.reset_index() if index else df.copy()
    for c in d.columns:
        if pd.api.types.is_datetime64_any_dtype(d[c]):
            d[c] = d[c].astype(str)
    ws.append([str(c) for c in d.columns])
    for row in d.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, 'item') else
                   (str(v) if not isinstance(v, (int, float, str, bool, type(None))) else v)) for v in row])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3864'); c.alignment = Alignment(wrap_text=True)
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(d.columns, 1):
        w = max([len(str(h))] + [len(str(x)) for x in d.iloc[:200, i - 1].tolist()]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(w, 10), 45)


sheet('events_up', PR['events_up'])
sheet('events_down', PR['events_down'])
ep = PR['lag_profile'].pivot(index='feature', columns='lag_days', values='median_percentile').add_prefix('pct_lag_').reset_index()
ep = ep.merge(PR['separation_up'], on='feature', how='left').sort_values('strength', ascending=False)
sheet('event_profiles', ep)
sheet('separation_down', PR['separation_down'])
sheet('events_by_year_sector', pd.concat([PR['up_by_year'].rename('up_events'), PR['down_by_year'].rename('down_events')], axis=1), index=True)
sheet('events_by_sector', pd.concat([PR['up_by_sector'].rename('up_events'), PR['down_by_sector'].rename('down_events')], axis=1).fillna(0), index=True)
cen = PR['clusters']['centroids'].copy(); cen.insert(0, 'cluster', range(len(cen)))
cen = cen.merge(PR['cluster_hit_rates'], on='cluster')
sheet('clusters', cen)
sheet('zec_like_and_rules', pd.DataFrame([{'item': 'zec_like_group (age>=730d, ATH<=-80%, base>=90d)', **PR['zec_like']}] +
                                          [{'item': f'rule_up[{i}]', 'feature': f, 'op': o, 'train_pos_quantile': q} for i, (f, o, q) in enumerate(PR['rule_up'])] +
                                          [{'item': f'rule_down[{i}]', 'feature': f, 'op': o, 'train_pos_quantile': q} for i, (f, o, q) in enumerate(PR['rule_down'])]))
sm = []
for col in ('feat_dist', 'dtw_dist', 'combined_rank'):
    t = SI[f'top20_{col}'].copy(); t.insert(0, 'method', col); sm.append(t.rename(columns={col: 'distance_or_rank'}))
sheet('similar_matches', pd.concat(sm, ignore_index=True))
sheet('similar_hit_curves', pd.concat([SI[f'curve_{c}'].assign(period='walk-forward') for c in ('feat_dist', 'dtw_dist', 'combined_rank')] +
                                      [SI[f'curve_post_{c}'].assign(period='post-dev 2024-07..2025-06') for c in ('feat_dist', 'dtw_dist', 'combined_rank')]))
mm = []
for tgt in ('up', 'down'):
    for per in ('metrics_wf', 'metrics_post_dev'):
        t = MO[tgt][per].copy(); t.insert(0, 'period', per); t.insert(0, 'target', tgt); mm.append(t)
MM = pd.concat(mm, ignore_index=True)
xs = pd.DataFrame([{'target': k[0], 'model': k[1], **v} for k, v in EV['xsec'].items()])
MM = MM.merge(xs.add_prefix('xsec_').rename(columns={'xsec_target': 'target', 'xsec_model': 'model'}), on=['target', 'model'], how='left')
sheet('model_metrics', MM)
fi = []
for tgt in ('up', 'down'):
    for f in MO[tgt]['info']:
        fi.append({'target': tgt, **{k: v for k, v in f.items() if not k.startswith('_') and k != 'rule_thresholds'}})
sheet('fold_info', pd.DataFrame(fi))
sheet('shap_up', MO['up']['shap'], index=True)
sheet('shap_down', MO['down']['shap'], index=True)

rows = []
for k, r in BT.items():
    if k.startswith('_'):
        continue
    for mode in ('native', 'unified'):
        for per in ('wf', 'post_dev'):
            rows.append({'variant': k, 'exit': mode, 'period': per, 'signals': r['n_signals'], **r[mode][per]})
for lab, key in (('LEGACY_FIXED long', 'L4_+M4_keep_NaN'), ('LEGACY_AS_IS long', 'L0_AS_IS')):
    for mode in ('native', 'unified'):
        rows.append({'variant': lab, 'exit': mode, 'period': 'wf', **LR[key][mode]['wf']})
for n in ('BTC_HOLD', 'EQUAL_WEIGHT'):
    rows.append({'variant': n, 'exit': '-', 'period': 'wf', **BR[n]['wf']})
sheet('backtest_metrics', pd.DataFrame(rows))
sheet('criteria_section11', EV['criteria'])
sheet('random_vs_new', pd.concat(EV['random'].values(), ignore_index=True))
sheet('bootstrap_vs_legacy', pd.DataFrame([{'variant': k[0], 'exit': k[1], 'ci_lo': v[0], 'ci_hi': v[1], 'diff': v[2]} for k, v in EV['bootstrap'].items()]))
tl = []
for k in ('NEW_XGB', 'NEW_BASELINE', 'NEW_LOGISTIC'):
    for mode in ('native', 'unified'):
        tl.append(BT[k][mode]['td'].assign(variant=k, exit=mode))
sheet('trade_log', pd.concat(tl, ignore_index=True))
yr = {}
for k, r in BT.items():
    if not k.startswith('_'):
        for mode in ('native', 'unified'):
            yr[f'{k} [{mode}]'] = r[mode]['yearly']
for mode in ('native', 'unified'):
    yr[f'LEGACY_FIXED long [{mode}]'] = LR['L4_+M4_keep_NaN'][mode]['yearly']
sheet('yearly_results', pd.DataFrame(yr).T, index=True)
rb = []
for k, r in BT['_robustness'].items():
    for mode in ('native', 'unified'):
        rb.append({'test': k, 'exit': mode, **{a: r[mode]['wf'].get(a) for a in ('trades', 'expectancy', 'expectancy_ex_top3', 'CAGR', 'MaxDD', 'Sharpe', 'Calmar')}})
for mode in ('native', 'unified'):
    rb.append({'test': 'BASE NEW_XGB', 'exit': mode, **{a: BT['NEW_XGB'][mode]['wf'].get(a) for a in ('trades', 'expectancy', 'expectancy_ex_top3', 'CAGR', 'MaxDD', 'Sharpe', 'Calmar')}})
    rb.append({'test': 'BASE NEW_XGB_regime', 'exit': mode, **{a: BT['NEW_XGB_regime'][mode]['wf'].get(a) for a in ('trades', 'expectancy', 'expectancy_ex_top3', 'CAGR', 'MaxDD', 'Sharpe', 'Calmar')}})
RB = pd.DataFrame(rb)
sheet('robustness', RB)
sheet('unit_tests', UT)
lp = []
for ex in ('native', 'unified'):
    c = LP[ex]['compare'].reset_index().rename(columns={'index': 'feature'}); c.insert(0, 'exit', ex); lp.append(c)
sheet('legacy_picks_profile', pd.concat(lp, ignore_index=True))
lb = []
for ex in ('native', 'unified'):
    for f, b in LP[ex]['buckets'].items():
        t = b.reset_index().rename(columns={b.index.name or 'index': 'bucket'}); t['bucket'] = t['bucket'].astype(str)
        t.insert(0, 'feature', f); t.insert(0, 'exit', ex); lb.append(t)
    s = LP[ex]['sector'].reset_index().rename(columns={'index': 'bucket', 'sector': 'bucket'}); s.insert(0, 'feature', 'sector'); s.insert(0, 'exit', ex); lb.append(s)
sheet('legacy_picks_buckets', pd.concat(lb, ignore_index=True))
out = os.path.join(ROOT, 'reports', f'new_walkforward_{DATE}.xlsx')
wb.save(out)
print('saved', out, '| unit tests:', UT['status'].value_counts().to_dict())
print(RB.round(4).to_string())
print(MO['up']['shap']['mean_abs_shap'].head(10).round(4).to_string())
print(MO['down']['shap']['mean_abs_shap'].head(10).round(4).to_string())
print(PR['up_by_year'].to_dict(), PR['down_by_year'].to_dict())
print(PR['up_by_sector'].head(8).to_dict())
print(PR['separation_up'].head(10)[['feature', 'event_median', 'other_median', 'auc', 'p_bonferroni']].round(4).to_string())
print(PR['cluster_hit_rates'].round(4).to_string()); print(PR['clusters']['centroids'].round(3).to_string())
for k in ('NEW_XGB', 'NEW_BASELINE_regime', 'NEW_LOGISTIC'):
    print(k, BT[k]['native']['yearly'].round(3).to_dict(), 'post-dev', {a: round(BT[k]['native']['post_dev'].get(a, np.nan), 4) for a in ('trades', 'expectancy', 'CAGR')})
print(SI['top20_feat_dist'].head(20).to_string())
