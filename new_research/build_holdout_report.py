"""reports/new_holdout_<date>.xlsx + charts from the one holdout run (no re-runs, no tuning)."""
import os
import pickle
import sys

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from new_research.holdout import CACHE, H0                    # noqa: E402

DATE = '20261018'
CH = os.path.join(ROOT, 'reports', 'charts')
R = pickle.load(open(os.path.join(CACHE, 'holdout_results.pkl'), 'rb'))
RES, THR = R['res'], R['thr']
EV = pickle.load(open(os.path.join(CACHE, 'holdout_eval.pkl'), 'rb'))
Z = pickle.load(open(os.path.join(CACHE, 'zec_test.pkl'), 'rb'))
RND = pickle.load(open(os.path.join(CACHE, 'holdout_random.pkl'), 'rb'))
S = pd.read_parquet(os.path.join(CACHE, 'scores_holdout.parquet')); S['day'] = pd.to_datetime(S['day'])


def reb(eq):
    e = eq[eq.index >= H0]
    return e / e.iloc[0]


# ---- supplementary (POST-HOC, labelled): false positives in the spec's own ZEC window Sep-Nov 2025
SPEC_W = (pd.Timestamp('2025-09-01'), pd.Timestamp('2025-11-30'))
supp = []
for k in ('NEW_LOGISTIC', 'NEW_BASELINE', 'LEGACY_FIXED_long', 'LEGACY_AS_IS_long'):
    sd = RES[k]['sig_df'].copy()
    sd['sday'] = pd.to_datetime(sd['ts']) - (pd.Timedelta(days=1) if k.startswith('NEW') else pd.Timedelta(hours=4))
    w = sd[(sd['sday'] >= SPEC_W[0]) & (sd['sday'] <= SPEC_W[1])]
    first = w.sort_values('ts').groupby('symbol').head(1)
    hit = np.nan
    if k.startswith('NEW'):
        lab = S.set_index(['symbol', 'day'])['up'].reindex(pd.MultiIndex.from_arrays([first['symbol'], first['sday'].dt.normalize()]))
        hit = (int(np.nansum(lab.values)), int(lab.notna().sum()))
    supp.append({'variant': k, 'window': 'POST-HOC 2025-09-01..2025-11-30 (spec text: ZEC ~55->650)', 'signals': len(w),
                 'distinct_coins': w['symbol'].nunique(), 'zec_signalled': bool((w['symbol'] == 'ZECUSDT').any()),
                 'distinct_coins_that_did_2x': hit[0] if isinstance(hit, tuple) else None,
                 'with_known_outcome': hit[1] if isinstance(hit, tuple) else None})
SUPP = pd.DataFrame(supp)

# ---- charts
fig, ax = plt.subplots(2, 1, figsize=(12, 9), sharex=True, gridspec_kw={'height_ratios': [2, 1]})
for k, ls in (('NEW_LOGISTIC', '-'), ('NEW_BASELINE', '-'), ('NEW_DOWN_BASELINE', '-'), ('NEW_DOWN_XGB', '-'),
              ('LEGACY_FIXED_long', ':'), ('LEGACY_FIXED_short', ':'), ('LEGACY_AS_IS_long', ':')):
    e = reb(RES[k]['native']['eq'])
    ax[0].plot(e.index, e.values, ls, label=k, lw=1.3)
    ax[1].plot(e.index, (e / e.cummax() - 1).values, ls, lw=1)
for k, col in (('BTC_HOLD', 'k'), ('EQUAL_WEIGHT', 'gray')):
    e = reb(RES['_bench'][k]['eq'])
    ax[0].plot(e.index, e.values, '--', color=col, label=k)
    ax[1].plot(e.index, (e / e.cummax() - 1).values, '--', color=col, lw=1)
ax[0].set_yscale('log'); ax[0].legend(fontsize=8); ax[0].grid(alpha=.3)
ax[0].set_title('FINAL HOLDOUT 2025-07-01 -> 2026-09-30 (one run, frozen models), native exit')
ax[1].set_title('Drawdown'); ax[1].grid(alpha=.3)
plt.tight_layout(); plt.savefig(os.path.join(CH, f'holdout_equity_drawdown_{DATE}.png'), dpi=120); plt.close()

zc = Z['zec_close']
fig, ax = plt.subplots(figsize=(13, 6))
ax.plot(zc.index, zc.values, color='k', lw=1.2, label='ZECUSDT daily close')
ax.set_yscale('log')
run = Z['run']
ax.axvline(run['low_day'], color='green', ls=':', label=f"low {run['low']:.1f} ({run['low_day'].date()})")
if run['x2_day'] is not None:
    ax.axvline(run['x2_day'], color='orange', ls=':', label=f"2x of low ({run['x2_day'].date()})")
marks = [('NEW_LOGISTIC', 'o', 'tab:blue'), ('NEW_BASELINE', '^', 'tab:cyan'), ('LEGACY_FIXED_long', 's', 'tab:red'),
         ('LEGACY_AS_IS_long', 'D', 'tab:pink'), ('NEW_DOWN_BASELINE', 'v', 'tab:purple'), ('LEGACY_FIXED_short', 'x', 'tab:brown')]
for i, (k, m, c) in enumerate(marks):
    sd = RES[k]['sig_df']
    sd = sd[sd['symbol'] == 'ZECUSDT']
    if len(sd):
        dd = pd.to_datetime(sd['ts']).dt.normalize()
        y = zc.reindex(dd).values * (1 + 0.06 * (i + 1))
        ax.scatter(dd, y, marker=m, s=14, color=c, label=f'{k} signals ({len(sd)})')
    td = RES[k]['native']['td']
    if len(td):
        t = td[td['symbol'] == 'ZECUSDT']
        ax.scatter(t['entry_time'], t['entry_px'], marker='*', s=200, color=c, edgecolor='k', zorder=5)
ax.set_title('ZEC in the holdout: signals of each variant (stars = trades actually opened by common_eval)')
ax.legend(fontsize=7, loc='upper left'); ax.grid(alpha=.3)
plt.tight_layout(); plt.savefig(os.path.join(CH, f'holdout_zec_signals_{DATE}.png'), dpi=120); plt.close()

fig, axs = plt.subplots(2, 4, figsize=(18, 8))
for row, k in enumerate(('NEW_LOGISTIC', 'NEW_DOWN_BASELINE')):
    for a, (mode, ex) in zip(axs[row], [('timing', 'native'), ('timing', 'unified'), ('rate', 'native'), ('rate', 'unified')]):
        rs = RND[k][(mode, ex)]['expectancy']
        a.hist(rs, bins=40, color='lightgray', edgecolor='gray')
        v = RES[k][ex]['wf']['expectancy']
        a.axvline(v, color='red', lw=2, label=f'{k} {v:.3f}'); a.axvline(np.percentile(rs, 95), color='blue', ls='--', label='P95')
        a.set_title(f'HOLDOUT {k} | RANDOM {mode} | {ex}', fontsize=8); a.legend(fontsize=7)
plt.tight_layout(); plt.savefig(os.path.join(CH, f'holdout_random_hist_{DATE}.png'), dpi=110); plt.close()

# ---- workbook
wb = Workbook(); wb.remove(wb.active)


def sheet(name, df, index=False):
    ws = wb.create_sheet(name[:31])
    d = df.reset_index() if index else df.copy()
    for c in d.columns:
        if pd.api.types.is_datetime64_any_dtype(d[c]):
            d[c] = d[c].astype(str)
    ws.append([str(c) for c in d.columns])
    for row in d.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, 'item') else
                   (str(v) if not isinstance(v, (int, float, str, bool, type(None))) else v)) for v in row])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3864'); c.alignment = Alignment(wrap_text=True)
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(d.columns, 1):
        w = max([len(str(h))] + [len(str(x)) for x in d.iloc[:200, i - 1].tolist()]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(w, 10), 45)


sheet('criteria_hypotheses', EV['criteria'])
sheet('all_variants', EV['all'])
sheet('random_holdout', EV['random'])
md = []
for k, v in EV['model'].items():
    md.append({'model': k, **(v if isinstance(v, dict) else {'value': v})})
sheet('model_metrics_holdout', pd.DataFrame(md))
sheet('zec_test', Z['table'])
sheet('zec_false_positives', Z['false_positives'])
sheet('zec_fp_posthoc_spec_window', SUPP)
zl = Z['reasons']['NEW_LOGISTIC_scores'].reset_index()
sheet('zec_new_logistic_scores', zl)
rr = []
for key in ('NEW_LOGISTIC_top5_at_best_day', 'NEW_LOGISTIC_top5_at_low'):
    if key in Z['reasons']:
        d, s = Z['reasons'][key]
        rr += [{'what': key, 'day': d, 'feature': f, 'contribution': float(v)} for f, v in s.items()]
if 'LEGACY_FIXED_long_best' in Z['reasons']:
    d, pr, s = Z['reasons']['LEGACY_FIXED_long_best']
    rr += [{'what': f'LEGACY_FIXED_long SHAP at highest-prob ZEC bar (prob {pr:.3f})', 'day': d, 'feature': f, 'contribution': float(v)} for f, v in s.items()]
sheet('zec_reasons', pd.DataFrame(rr))
sheet('zec_rule_conditions', Z['reasons']['NEW_BASELINE_rule_on_zec'])
for k in ('NEW_LOGISTIC', 'NEW_DOWN_BASELINE', 'NEW_DOWN_XGB', 'LEGACY_FIXED_long', 'LEGACY_FIXED_short'):
    sheet(f'trades_{k}'[:31], pd.concat([RES[k][m]['td'].assign(exit=m) for m in ('native', 'unified')], ignore_index=True))
yr = {f'{k} [{m}]': RES[k][m]['yearly'] for k in RES if not k.startswith('_') for m in ('native', 'unified')}
sheet('yearly_holdout', pd.DataFrame(yr).T, index=True)
out = os.path.join(ROOT, 'reports', f'new_holdout_{DATE}.xlsx')
wb.save(out)
print('saved', out)
print(SUPP.to_string())
